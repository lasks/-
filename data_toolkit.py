# -*- coding: utf-8 -*-
"""
数据处理工具箱 —— 现场实施最常用的几个功能都在这。
用法：
  python data_toolkit.py check --source xxx.csv --target yyy.csv --key id
  python data_toolkit.py import --csv data.csv --db config.json --table orders
  python data_toolkit.py dedupe --input raw.csv --output clean.csv --key id,name

依赖：pip install pymysql openpyxl
"""

import os, sys, csv, json, hashlib, logging
from datetime import datetime
from pathlib import Path

# ── 日志 ──────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler()]
)
log = logging.getLogger("toolkit")


# ============================================================
# 1. 文件读写 —— 自动处理编码
# ============================================================

def read_csv(filepath):
    """读CSV，GBK/UTF-8 自动试"""
    for enc in ["utf-8", "gbk", "gb2312", "utf-8-sig"]:
        try:
            with open(filepath, "r", encoding=enc) as f:
                reader = csv.DictReader(f)
                rows = list(reader)
                log.info(f"读取 {filepath} [{enc}] : {len(rows)} 行")
                return rows
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise Exception(f"读不了 {filepath}，编码都不对")


def write_csv(data, filepath):
    """写CSV，带 BOM 方便 Excel 打开"""
    if not data:
        log.warning("数据是空的，不写了")
        return
    headers = list(data[0].keys())
    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(data)
    log.info(f"写入 {filepath}: {len(data)} 行")


def read_excel(filepath, sheet=0):
    """读Excel，返回字典列表"""
    from openpyxl import load_workbook
    wb = load_workbook(filepath, read_only=True)
    ws = wb[sheet] if isinstance(sheet, str) else wb.worksheets[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [str(h or f"col_{i}") for i, h in enumerate(rows[0])]
    result = [{headers[i]: v for i, v in enumerate(r)} for r in rows[1:]]
    log.info(f"读取 {filepath}: {len(result)} 行")
    return result


# ============================================================
# 2. 数据库操作 —— 不想用ORM，直接写SQL
# ============================================================

class DB:
    """简单的MySQL操作封装"""

    def __init__(self, config):
        import pymysql
        self.conn = pymysql.connect(
            host=config["host"],
            port=config.get("port", 3306),
            user=config["user"],
            password=config["password"],
            database=config["database"],
            charset="utf8mb4",
            connect_timeout=10,
        )
        log.info(f"连上数据库: {config['host']}/{config['database']}")

    def query(self, sql, params=None):
        """查数据，返回字典列表"""
        with self.conn.cursor() as cur:
            cur.execute(sql, params)
            cols = [d[0] for d in cur.description]
            return [dict(zip(cols, r)) for r in cur.fetchall()]

    def execute(self, sql, params=None):
        """增删改，返回影响行数"""
        with self.conn.cursor() as cur:
            n = cur.execute(sql, params)
            self.conn.commit()
            return n

    def close(self):
        self.conn.close()
        log.info("数据库连接关了")


# ============================================================
# 3. 数据校验 —— 现场80%的工作是这个
# ============================================================

def check_count(source_count, target_count, label=""):
    """比条数"""
    diff = source_count - target_count
    pct = abs(diff) / max(source_count, 1) * 100
    r = {"label": label, "source": source_count, "target": target_count,
         "diff": diff, "diff_pct": round(pct, 2)}
    if diff == 0:
        log.info(f"[{label}] 条数一致: {source_count}")
    else:
        log.warning(f"[{label}] 源={source_count} 目标={target_count} 差{diff}条({pct:.2f}%)")
    return r


def find_missing(source_file, target_file, key_field):
    """找出源有但目标没有的记录"""
    src = read_csv(source_file)
    tgt = read_csv(target_file)
    src_keys = {r[key_field] for r in src if r.get(key_field)}
    tgt_keys = {r[key_field] for r in tgt if r.get(key_field)}
    missing = src_keys - tgt_keys
    extra = tgt_keys - src_keys
    log.info(f"源:{len(src_keys)} 目标:{len(tgt_keys)} 目标缺:{len(missing)} 多余:{len(extra)}")
    if missing:
        log.warning(f"目标缺失样例(前10): {list(missing)[:10]}")
    return {"missing": list(missing), "extra": list(extra)}


def check_nulls(filepath, column):
    """检查某列的空值率"""
    rows = read_csv(filepath)
    total = len(rows)
    nulls = sum(1 for r in rows if not r.get(column) or str(r[column]).strip() == "")
    rate = nulls / max(total, 1) * 100
    log.info(f"[{column}] 空值率: {nulls}/{total} = {rate:.2f}%")
    return {"column": column, "total": total, "nulls": nulls, "rate": round(rate, 2)}


# ============================================================
# 4. 数据处理
# ============================================================

def dedupe(input_file, output_file, key_columns):
    """按指定列去重，保留第一条"""
    rows = read_csv(input_file)
    seen = set()
    result = []
    dupes = 0
    for r in rows:
        k = tuple(str(r.get(c, "")) for c in key_columns)
        if k in seen:
            dupes += 1
            continue
        seen.add(k)
        result.append(r)
    write_csv(result, output_file)
    log.info(f"去重: {len(rows)}→{len(result)}条, 去掉{dupes}条重复")
    return len(result)


def merge_csv(file_list, output_file, on_column):
    """多个CSV按某列合并（left join）"""
    if len(file_list) < 2:
        log.error("至少两个文件")
        return
    base = read_csv(file_list[0])
    base_map = {r[on_column]: r for r in base if r.get(on_column)}
    for fp in file_list[1:]:
        extra = read_csv(fp)
        for r in extra:
            k = r.get(on_column)
            if k and k in base_map:
                base_map[k].update(r)
    result = list(base_map.values())
    write_csv(result, output_file)
    log.info(f"合并完成: {len(result)} 条 → {output_file}")


def csv_to_db(csv_path, db_config, table, truncate=False):
    """CSV导入MySQL"""
    rows = read_csv(csv_path)
    if not rows:
        return 0
    db = DB(db_config)
    if truncate:
        db.execute(f"TRUNCATE TABLE {table}")

    headers = list(rows[0].keys())
    placeholders = ", ".join(["%s"] * len(headers))
    cols = ", ".join([f"`{h}`" for h in headers])
    sql = f"INSERT INTO {table} ({cols}) VALUES ({placeholders})"

    batch = 500
    total = 0
    for i in range(0, len(rows), batch):
        chunk = rows[i:i + batch]
        vals = [[r.get(h) for h in headers] for r in chunk]
        with db.conn.cursor() as cur:
            cur.executemany(sql, vals)
            db.conn.commit()
        total += len(chunk)
        log.info(f"  导入 {total}/{len(rows)}")
    db.close()
    log.info(f"导入完成: {total} 条 → {table}")
    return total


# ============================================================
# 5. 文件完整性
# ============================================================

def file_md5(filepath):
    """算文件MD5"""
    h = hashlib.md5()
    with open(filepath, "rb") as f:
        while chunk := f.read(8192):
            h.update(chunk)
    return h.hexdigest()


def verify_dir(source_dir, target_dir):
    """比对两个目录的文件是否一致"""
    import glob
    src_files = {os.path.basename(f): f for f in glob.glob(os.path.join(source_dir, "*"))}
    tgt_files = {os.path.basename(f): f for f in glob.glob(os.path.join(target_dir, "*"))}
    all_names = set(src_files.keys()) | set(tgt_files.keys())

    ok, fail = 0, 0
    for name in sorted(all_names):
        if name not in src_files:
            log.warning(f"[缺] 源端没有: {name}")
            fail += 1
        elif name not in tgt_files:
            log.warning(f"[缺] 目标端没有: {name}")
            fail += 1
        else:
            s_md5 = file_md5(src_files[name])
            t_md5 = file_md5(tgt_files[name])
            if s_md5 == t_md5:
                ok += 1
            else:
                log.warning(f"[不一致] {name} src={s_md5[:8]} tgt={t_md5[:8]}")
                fail += 1
    log.info(f"校验结果: {ok} OK / {fail} 异常")
    return {"ok": ok, "fail": fail}


# ============================================================
# 6. 命令行入口
# ============================================================

if __name__ == "__main__":
    import argparse

    p = argparse.ArgumentParser(description="现场数据处理工具箱")
    sub = p.add_subparsers(dest="cmd")

    # check: 比对两个CSV的条数和缺失
    c = sub.add_parser("check", help="比对两个CSV")
    c.add_argument("--source", required=True)
    c.add_argument("--target", required=True)
    c.add_argument("--key", required=True, help="比对用的主键字段")

    # import: CSV导入数据库
    imp = sub.add_parser("import", help="CSV导入MySQL")
    imp.add_argument("--csv", required=True)
    imp.add_argument("--db-config", required=True, help="数据库配置JSON文件")
    imp.add_argument("--table", required=True)
    imp.add_argument("--truncate", action="store_true", help="先清空表")

    # dedupe: 去重
    d = sub.add_parser("dedupe", help="CSV去重")
    d.add_argument("--input", required=True)
    d.add_argument("--output", required=True)
    d.add_argument("--key", required=True, help="去重依据的列名，多个用逗号分隔")

    # nulls: 空值检查
    n = sub.add_parser("nulls", help="检查空值率")
    n.add_argument("--file", required=True)
    n.add_argument("--column", required=True)

    # merge: 合并多个CSV
    m = sub.add_parser("merge", help="合并多个CSV")
    m.add_argument("--files", nargs="+", required=True)
    m.add_argument("--output", required=True)
    m.add_argument("--on", required=True, dest="on_column", help="合并用的列名")

    # verify: 文件完整性校验
    v = sub.add_parser("verify", help="比对两个目录的文件")
    v.add_argument("--src", required=True)
    v.add_argument("--tgt", required=True)

    args = p.parse_args()

    if args.cmd == "check":
        find_missing(args.source, args.target, args.key)

    elif args.cmd == "import":
        with open(args.db_config, "r") as f:
            cfg = json.load(f)
        csv_to_db(args.csv, cfg, args.table, args.truncate)

    elif args.cmd == "dedupe":
        keys = [k.strip() for k in args.key.split(",")]
        dedupe(args.input, args.output, keys)

    elif args.cmd == "nulls":
        check_nulls(args.file, args.column)

    elif args.cmd == "merge":
        merge_csv(args.files, args.output, args.on_column)

    elif args.cmd == "verify":
        verify_dir(args.src, args.tgt)

    else:
        p.print_help()
